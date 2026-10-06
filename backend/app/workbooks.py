"""XLSX workbook ingestion: worksheet listing and one-sheet normalization.

An .xlsx workbook is a ZIP package of XML parts, read here with openpyxl in
read-only mode so that large worksheets stream through instead of loading whole.
The chosen worksheet is normalized into CSV text and then analyzed by the same
engine as a CSV file, so both input formats follow one set of Missing Value
rules and one definition of an Input Row.

Excel's file format defines a maximum of 1,048,576 worksheet rows, header row
included. The application reads every row a worksheet actually holds and never
truncates it: a worksheet at the limit is counted to its full 1,048,575 data
rows.
"""

import csv
import datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

EXCEL_WORKSHEET_ROW_LIMIT = 1_048_576

WORKSHEET_CSV_NAME = "worksheet.csv"

INVALID_WORKBOOK_MESSAGE = "The uploaded file could not be read as an XLSX workbook."
NO_WORKSHEETS_MESSAGE = "The uploaded workbook contains no worksheets."
CHOOSE_WORKSHEET_MESSAGE = "Choose exactly one worksheet of the workbook to analyze."
CSV_HAS_NO_WORKSHEETS_MESSAGE = (
    "Worksheet selection applies to .xlsx workbooks; a CSV file has no worksheets."
)


class WorkbookFileError(ValueError):
    """Raised when an uploaded file cannot be read as an XLSX workbook."""


class WorksheetSelectionError(ValueError):
    """Raised when the chosen worksheet cannot be analyzed."""


def is_workbook_file_name(file_name: str) -> bool:
    return Path(file_name).suffix.lower() == ".xlsx"


def list_worksheet_names(workbook_path: Path) -> list[str]:
    workbook = open_workbook(workbook_path)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def write_worksheet_to_csv(workbook_path: Path, worksheet_name: str, csv_path: Path) -> Path:
    """Write the chosen worksheet to CSV text and return the written path.

    Every field is quoted so that values with spaces, quotes, or line breaks
    reach the analysis engine unchanged.
    """
    workbook = open_workbook(workbook_path)
    try:
        worksheet = require_worksheet(workbook, worksheet_name)
        rows_with_values = write_worksheet_rows_to_csv(worksheet, csv_path)
    finally:
        workbook.close()

    if rows_with_values == 0:
        csv_path.unlink(missing_ok=True)
        raise WorksheetSelectionError(f"The selected worksheet {worksheet_name!r} is empty.")

    return csv_path


def write_worksheet_rows_to_csv(worksheet, csv_path: Path) -> int:
    """Write every worksheet row to CSV text and return how many rows carry values.

    A row where every cell is empty becomes an empty line, so the engine applies
    the same blank-line rule as for a CSV file: not an Input Row, except in a
    one-column file where an empty line is one Input Row with a missing value.
    """
    rows_with_values = 0
    with csv_path.open("w", encoding="utf-8", newline="") as csv_file:
        writer = csv.writer(csv_file, quoting=csv.QUOTE_ALL)
        for worksheet_row in worksheet.iter_rows(values_only=True):
            if worksheet_row_is_blank(worksheet_row):
                writer.writerow([])
                continue
            writer.writerow([worksheet_cell_text(cell_value) for cell_value in worksheet_row])
            rows_with_values = rows_with_values + 1

    return rows_with_values


def worksheet_row_is_blank(worksheet_row: tuple[object, ...]) -> bool:
    return all(cell_value is None for cell_value in worksheet_row)


def worksheet_cell_text(cell_value: object) -> str:
    """Render one worksheet cell as the text the analysis engine reads.

    Formula cells are read as their formula text: the application never
    evaluates formulas, so a formula is present text like any other.
    """
    if cell_value is None:
        return ""
    if isinstance(cell_value, bool):
        return "TRUE" if cell_value else "FALSE"
    if isinstance(cell_value, datetime.datetime):
        return cell_value.isoformat(sep=" ")
    if isinstance(cell_value, (datetime.date, datetime.time)):
        return cell_value.isoformat()
    return str(cell_value)


def open_workbook(workbook_path: Path):
    """Open a workbook for streaming reads, rejecting unreadable files clearly."""
    try:
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    except (InvalidFileException, BadZipFile, KeyError, ParseError, OSError) as error:
        raise WorkbookFileError(INVALID_WORKBOOK_MESSAGE) from error

    if not workbook.sheetnames:
        workbook.close()
        raise WorkbookFileError(NO_WORKSHEETS_MESSAGE)

    return workbook


def require_worksheet(workbook, worksheet_name: str):
    """Return exactly the named worksheet, listing the workbook's worksheets otherwise."""
    if worksheet_name not in workbook.sheetnames:
        available_worksheets = ", ".join(workbook.sheetnames)
        raise WorksheetSelectionError(
            f"The workbook has no worksheet named {worksheet_name!r}. "
            f"Worksheets in this workbook: {available_worksheets}."
        )

    return workbook[worksheet_name]
